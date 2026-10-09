"""Testes do conferir.py (sem IA) com perfil, vaga e currículo fictícios."""
import shutil
import sys
import tempfile
import time
import unittest
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
import conferir  # noqa: E402

PERFIL = """# Meu perfil de carreira
## Experiência
### Analista de Dados — Empresa Alfa (03/2021–atual)
- Automatizei o relatório mensal, reduzindo o fechamento de 10 para 6 dias
- Painéis em Power BI usados por 1.500 pessoas; satisfação subiu 30%
### Assistente de Dados — Empresa Beta (01/2018–06/2020)
- Consultas em SQL para a área comercial
## Habilidades e ferramentas (com nível)
- SQL — avançado
- Power BI — avançado
- Jira — intermediário
"""
VAGA = """Analista de Dados Pleno
Responsabilidades:
- Criar painéis para a diretoria
Requisitos:
- Experiência com SQL avançado
- Conhecimento em Tableau
- Python para análise de dados
- Inglês fluente
Benefícios:
- Vale refeição
"""


def cv(**mudancas):
    base = {
        "nome": "Ana Exemplo", "contato": ["ana@exemplo.com"], "titulo": "Analista de Dados",
        "resumo": "Analista de dados com painéis usados por 1,5 mil pessoas.",
        "competencias": [{"grupo": "Dados", "itens": ["SQL", "Power BI"]}],
        "experiencia": [
            {"empresa": "Empresa Alfa", "periodo": "03/2021 – atual", "cargos": [
                {"cargo": "Analista de Dados", "itens": ["Reduziu o fechamento de 10 para 6 dias automatizando o relatório",
                                                         "Satisfação subiu 30 % com os novos painéis"]}]},
            {"empresa": "Empresa Beta", "periodo": "01/2018 – 06/2020", "cargos": [
                {"cargo": "Assistente de Dados", "itens": ["Consultas em SQL para a área comercial"]}]}],
    }
    base.update(mudancas)
    return base


class TestFatos(unittest.TestCase):
    def test_fiel_ao_perfil_nao_aponta(self):
        self.assertEqual(conferir.conferir_fatos(cv(), PERFIL), [])

    def test_numero_inventado_bloqueia(self):
        c = cv()
        c["experiencia"][0]["cargos"][0]["itens"].append("Aumentou a receita em 45%")
        r = conferir.conferir(c, PERFIL)
        self.assertEqual(r["veredito"], "bloquear")
        self.assertTrue(any(p["tipo"] == "numero" and "45%" in p["valor"] for p in r["fatos"]))

    def test_numeros_escritos_de_outro_jeito(self):
        for texto in ("painéis para 1500 pessoas", "painéis para 1.500 pessoas", "satisfação subiu trinta por cento"):
            self.assertEqual(conferir.conferir_fatos(cv(resumo=texto), PERFIL), [], texto)
        self.assertEqual(len(conferir.conferir_fatos(cv(resumo="painéis para 2 mil pessoas"), PERFIL)), 1)

    def test_empresa_cargo_e_data(self):
        c = cv()
        c["experiencia"].append({"empresa": "Empresa Gama", "periodo": "02/2015 – 12/2016",
                                 "cargos": [{"cargo": "Coordenador", "itens": []}]})
        tipos = {p["tipo"] for p in conferir.conferir_fatos(c, PERFIL)}
        self.assertEqual(tipos, {"empresa", "cargo", "data"})
        self.assertEqual(conferir.conferir(c, PERFIL)["veredito"], "conferir")


class TestRequisitos(unittest.TestCase):
    def test_tem_sustentado_lacuna_e_cobertura(self):
        reqs, cob, lacunas = conferir.conferir_requisitos(cv(), PERFIL, VAGA)
        sit = {r["texto"]: r["situacao"] for r in reqs}
        self.assertEqual(sit["Experiência com SQL avançado"], "tem")
        self.assertEqual(sit["Conhecimento em Tableau"], "sustentado")  # Power BI é da mesma família
        self.assertEqual(sit["Python para análise de dados"], "sustentado")  # "dados" está no perfil
        self.assertEqual(sit["Inglês fluente"], "lacuna")
        self.assertNotIn("Vale refeição", sit)
        self.assertNotIn("Criar painéis para a diretoria", sit)
        self.assertEqual(lacunas, [])
        self.assertGreater(cob["total"], 0)
        self.assertIn("python", cob["faltando"])

    def test_lacuna_como_competencia_bloqueia(self):
        c = cv(competencias=[{"grupo": "Idiomas", "itens": ["Inglês fluente"]}])
        r = conferir.conferir(c, PERFIL, VAGA)
        self.assertEqual(r["veredito"], "bloquear")
        self.assertTrue(any(p["tipo"] == "lacuna" for p in r["fatos"]))

    def test_vaga_sem_cabecalho_usa_marcadores(self):
        reqs = conferir.requisitos_da_vaga("Venha trabalhar conosco.\n- SQL\n- Excel avançado\n")
        self.assertEqual(reqs, ["SQL", "Excel avançado"])


class TestATS(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_tabela_e_cabecalho(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Ana Exemplo")
        doc.add_table(rows=1, cols=2)
        doc.sections[0].header.paragraphs[0].text = "ana@exemplo.com"
        caminho = self.tmp / "cv.docx"
        doc.save(caminho)
        msgs = " ".join(p["mensagem"] for p in conferir.conferir_ats(cv(), caminho))
        self.assertIn("tabela", msgs)
        self.assertIn("cabeçalho", msgs)

    def test_docx_do_motor_passa_e_dados_pessoais_no_us(self):
        import curriculo
        c = cv(tecnicas={"formato": "us"}, contato=["ana@exemplo.com", "Married, 35 years old"])
        caminho = self.tmp / "cv.docx"
        curriculo.Montador(c).montar().save(caminho)
        pontos = conferir.conferir_ats(c, caminho)
        self.assertEqual([p["nivel"] for p in pontos], ["bloquear"])
        self.assertEqual(conferir.conferir_ats(cv(), caminho), [])

    def test_rapido(self):
        inicio = time.time()
        conferir.conferir(cv(), PERFIL * 20, VAGA * 10)
        self.assertLess(time.time() - inicio, 5)


if __name__ == "__main__":
    unittest.main()
