"""Testes do primeiros_passos.py (python -m unittest discover -s tests). Materiais fictícios gerados aqui."""
import io
import json
import shutil
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import ia  # noqa: E402
import primeiros_passos as pp  # noqa: E402

CPF = "123.456.789-09"


def pdf_minimo(linhas):
    """PDF de uma página com texto (Helvetica, só ASCII); sem linhas, um PDF sem texto (como um digitalizado)."""
    conteudo = ("BT /F1 12 Tf 72 720 Td " + " ".join(f"({l}) Tj 0 -16 Td" for l in linhas) + " ET") if linhas else ""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> "
            "/Contents 4 0 R >>",
            f"<< /Length {len(conteudo)} >>\nstream\n{conteudo}\nendstream",
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    saida, posicoes = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        posicoes.append(len(saida))
        saida += f"{i} 0 obj\n{o}\nendobj\n".encode("latin-1")
    xref = len(saida)
    saida += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    saida += "".join(f"{p:010d} 00000 n \n" for p in posicoes).encode()
    saida += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return saida


def docx_bytes(paragrafos, tabela=None):
    import docx
    d = docx.Document()
    for p in paragrafos:
        d.add_paragraph(p)
    if tabela:
        t = d.add_table(rows=len(tabela), cols=len(tabela[0]))
        for i, linha in enumerate(tabela):
            for j, c in enumerate(linha):
                t.cell(i, j).text = c
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def zip_linkedin(completo=False):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("Profile.csv", "﻿First Name,Last Name,Headline,Summary,Geo Location\n"
                                  "Ana,Exemplo,Analista de Dados,Trabalho com dados há 5 anos.,\"Curitiba, PR\"\n")
        z.writestr("Positions.csv", "Company Name,Title,Description,Location,Started On,Finished On\n"
                                    "Empresa Alfa,Analista de Dados,\"Painéis de vendas\nAutomação de relatórios\","
                                    "Curitiba,Mar 2021,\n"
                                    "Empresa Beta,Assistente de Dados,,Curitiba,Jan 2018,Jun 2020\n")
        z.writestr("Skills.csv", "Name\nSQL\nPower BI\n")
        if completo:
            z.writestr("Education.csv", "School Name,Start Date,End Date,Notes,Degree Name,Activities\n"
                                        "Universidade Exemplo,2014,2018,,Bacharelado em Estatística,\n")
            z.writestr("Certifications.csv", "Name,Url,Authority,Started On,Finished On,License Number\n"
                                             "Certificação Exemplo,,Instituto Exemplo,Mai 2022,,\n")
            z.writestr("Languages.csv", "Name,Proficiency\nInglês,Profissional\n")
    return buf.getvalue()


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.patches = [mock.patch.object(pp, "ANEXOS", self.tmp / "anexos" / "primeiros-passos"),
                        mock.patch.object(pp, "ANTERIORES", self.tmp / "anexos" / "perfis-anteriores"),
                        mock.patch.object(pp, "PROGRESSO", self.tmp / ".cache" / "primeiros-passos.json"),
                        mock.patch.object(pp, "RAIZ", self.tmp),
                        mock.patch.object(pp.ocr, "disponivel", return_value=False)]  # sem OCR real nos testes
        for p in self.patches:
            p.start()
        shutil.copy(Path(__file__).resolve().parent.parent / "perfil.exemplo.md", self.tmp / "perfil.exemplo.md")
        self.cfg = {"perfil": str(self.tmp / "perfil.md")}

    def tearDown(self):
        for p in self.patches:
            p.stop()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def arquivo(self, nome, dados):
        caminho = self.tmp / nome
        caminho.write_bytes(dados)
        return caminho


class TestPrivacidade(unittest.TestCase):
    def test_mascara_documentos(self):
        texto = f"CPF: {CPF}\nRG: 12.345.678-9\nData de nascimento: 01/02/1990\nCPF sem pontos 12345678909"
        limpo = pp.mascarar_documentos(texto)
        for proibido in (CPF, "12.345.678-9", "01/02/1990", "12345678909"):
            self.assertNotIn(proibido, limpo)
        self.assertIn("[removido]", limpo)

    def test_nao_mascara_telefone_nem_datas_de_cargo(self):
        texto = "Telefone (41) 99999-8888 — Analista (03/2021–atual)"
        self.assertEqual(pp.mascarar_documentos(texto), texto)

    def test_reconhece_contato(self):
        c = pp.reconhecer_contato("ana@exemplo.com.br | (41) 99999-8888 | linkedin.com/in/ana-exemplo")
        self.assertEqual(c["emails"], ["ana@exemplo.com.br"])
        self.assertEqual(c["telefones"], ["(41) 99999-8888"])
        self.assertEqual(c["linkedin"], "linkedin.com/in/ana-exemplo")


class TestExtracao(Base):
    def test_docx_com_tabela_e_cpf(self):
        dados = docx_bytes(["Ana Exemplo", f"CPF {CPF}", "ana@exemplo.com"], [["SQL", "Avançado"]])
        r = pp.extrair(self.arquivo("cv.docx", dados))
        self.assertEqual(r["tipo"], "curriculo_docx")
        self.assertIn("Ana Exemplo", r["texto"])
        self.assertIn("SQL | Avançado", r["texto"])
        self.assertNotIn(CPF, r["texto"])
        self.assertEqual(r["contato"]["emails"], ["ana@exemplo.com"])

    def test_pdf(self):
        r = pp.extrair(self.arquivo("cv.pdf", pdf_minimo(["Ana Exemplo", "Analista de Dados", f"CPF {CPF}"])))
        self.assertEqual(r["tipo"], "curriculo_pdf")
        self.assertIn("Analista de Dados", r["texto"])
        self.assertNotIn(CPF, r["texto"])

    def test_pdf_do_linkedin(self):
        r = pp.extrair(self.arquivo("Profile.pdf", pdf_minimo(["www.linkedin.com/in/ana-exemplo", "Experience"])))
        self.assertEqual(r["tipo"], "linkedin_pdf")

    def test_pdf_sem_texto_avisa(self):
        r = pp.extrair(self.arquivo("scan.pdf", pdf_minimo([])))
        self.assertEqual(r["texto"], "")
        self.assertTrue(any("imagem" in a for a in r["avisos"]))

    def test_zip_do_linkedin(self):
        r = pp.extrair(self.arquivo("linkedin.zip", zip_linkedin()))
        self.assertEqual(r["tipo"], "linkedin_zip")
        c = r["campos"]
        self.assertEqual(c["perfil"]["titulo"], "Analista de Dados")
        self.assertEqual([x["empresa"] for x in c["cargos"]], ["Empresa Alfa", "Empresa Beta"])
        self.assertEqual(c["competencias"], ["SQL", "Power BI"])
        self.assertTrue(any("Education.csv" in a for a in r["avisos"]))
        completo = pp.extrair(self.arquivo("linkedin2.zip", zip_linkedin(completo=True)))
        self.assertEqual(completo["avisos"], [])
        self.assertEqual(completo["campos"]["idiomas"], [{"idioma": "Inglês", "nivel": "Profissional"}])

    def test_formatos_recusados(self):
        with self.assertRaises(pp.PerfilErro):
            pp.extrair(self.arquivo("antigo.doc", b"x"))
        with self.assertRaises(pp.PerfilErro):
            pp.extrair(self.arquivo("foto.png", b"x"))
        z = io.BytesIO()
        with zipfile.ZipFile(z, "w") as arq:
            arq.writestr("outra.txt", "nada")
        with self.assertRaises(pp.PerfilErro):
            pp.extrair(self.arquivo("outro.zip", z.getvalue()))
        with self.assertRaises(pp.PerfilErro):
            pp.extrair(self.arquivo("quebrado.zip", b"PK nada"))


CERTIFICADO = ["CERTIFICADO DE PARTICIPACAO", "Certificamos que Ana Exemplo", "concluiu o curso Introducao ao SQL",
               "Carga horaria: 12 horas"]


class TestAreas(Base):
    """Spec 037: materiais em três áreas (LinkedIn, currículo, documento), Mover para e o aviso de área trocada."""

    def test_area_de(self):
        esperado = {"linkedin_pdf": "linkedin", "linkedin_zip": "linkedin", "curriculo_pdf": "curriculo",
                    "curriculo_docx": "curriculo", "texto": "curriculo", "documento_pdf": "documento",
                    "documento_docx": "documento"}
        self.assertEqual({t: pp.area_de(t) for t in esperado}, esperado)

    def test_extrair_com_area(self):
        pdf = self.arquivo("cert.pdf", pdf_minimo(CERTIFICADO))
        self.assertEqual(pp.extrair(pdf, area="documento")["tipo"], "documento_pdf")
        self.assertEqual(pp.extrair(pdf, area="linkedin")["tipo"], "linkedin_pdf")  # vale a área escolhida
        self.assertEqual(pp.extrair(pdf)["tipo"], "curriculo_pdf")  # sem área, a adivinhação de hoje
        docx = self.arquivo("carta.docx", docx_bytes(["Carta de recomendação"]))
        self.assertEqual(pp.extrair(docx, area="documento")["tipo"], "documento_docx")
        li = self.arquivo("Profile.pdf", pdf_minimo(["www.linkedin.com/in/ana-exemplo", "Experience"]))
        self.assertEqual(pp.extrair(li, area="curriculo")["tipo"], "curriculo_pdf")

    def test_formatos_por_area(self):
        z = self.arquivo("linkedin.zip", zip_linkedin())
        with self.assertRaisesRegex(pp.PerfilErro, "ZIP do LinkedIn vai na área do LinkedIn"):
            pp.extrair(z, area="curriculo")
        with self.assertRaisesRegex(pp.PerfilErro, "ZIP do LinkedIn vai na área do LinkedIn"):
            pp.extrair(z, area="documento")
        self.assertEqual(pp.extrair(z, area="linkedin")["tipo"], "linkedin_zip")
        with self.assertRaisesRegex(pp.PerfilErro, "PDF do perfil do LinkedIn ou o ZIP"):
            pp.extrair(self.arquivo("cv.docx", docx_bytes(["Ana"])), area="linkedin")
        with self.assertRaisesRegex(pp.PerfilErro, "área inválida"):
            pp.extrair(z, area="x")
        with self.assertRaisesRegex(pp.PerfilErro, ".doc antigo"):
            pp.extrair(self.arquivo("antigo.doc", b"x"), area="documento")

    def test_adicionar_e_mover(self):
        e = pp.estado_vazio()
        doc = pp.adicionar_material(e, "cert.pdf", pdf_minimo(CERTIFICADO), area="documento")
        self.assertEqual(doc["tipo"], "documento_pdf")
        with self.assertRaises(pp.PerfilErro):
            pp.adicionar_material(e, "cv.docx", docx_bytes(["Ana"]), area="linkedin")
        self.assertEqual(len(e["materiais"]), 1)
        self.assertFalse(list(pp.ANEXOS.glob("*.docx")))  # recusado antes de gravar o arquivo
        cv = pp.adicionar_material(e, "cv.docx", docx_bytes(["Ana"]))
        antes = dict(cv)
        movido = pp.mover_material(e, cv["id"], "documento")
        self.assertEqual(movido["tipo"], "documento_docx")
        self.assertEqual({k: movido[k] for k in ("id", "arquivo", "texto")}, {k: antes[k] for k in ("id", "arquivo", "texto")})
        self.assertEqual(pp.mover_material(e, doc["id"], "linkedin")["tipo"], "linkedin_pdf")
        with self.assertRaisesRegex(pp.PerfilErro, "PDF do perfil do LinkedIn ou o ZIP"):
            pp.mover_material(e, cv["id"], "linkedin")  # Word não vai para o LinkedIn
        colado = pp.adicionar_material(e, "", texto="Sou analista.")
        with self.assertRaisesRegex(pp.PerfilErro, "texto colado fica em Currículos"):
            pp.mover_material(e, colado["id"], "documento")
        with self.assertRaises(pp.PerfilErro):
            pp.mover_material(e, "m99", "documento")
        with self.assertRaisesRegex(pp.PerfilErro, "área inválida"):
            pp.mover_material(e, cv["id"], None)

    def test_sugestao(self):
        sug = lambda tipo, texto, arquivo="anexos/a.pdf": (pp.sugestao_de_area(
            {"tipo": tipo, "texto": texto, "arquivo": arquivo}) or {}).get("area")
        cert = "\n".join(CERTIFICADO)
        li = "Ana Exemplo\nwww.linkedin.com/in/ana-exemplo\nExperience\nAnalista"
        cv_longo = "Ana Exemplo\nExperiência\n" + "Analista de dados, painéis e relatórios. " * 100 + "\nCertificado em SQL"
        self.assertEqual(sug("curriculo_pdf", cert), "documento")
        self.assertEqual(sug("curriculo_pdf", "CERTIFICAMOS que Ana concluiu. Carga Horária: 8 horas"), "documento")
        self.assertEqual(sug("curriculo_pdf", "Certificate of Completion — Ana Example"), "documento")
        self.assertIsNone(sug("curriculo_pdf", cv_longo))  # currículo longo que cita "certificado"
        self.assertEqual(sug("curriculo_pdf", li), "linkedin")
        self.assertEqual(sug("documento_pdf", li), "linkedin")
        self.assertIsNone(sug("documento_pdf", cert))
        self.assertIsNone(sug("documento_docx", li, "anexos/a.docx"))  # Word não vai para o LinkedIn
        self.assertEqual(sug("linkedin_pdf", cert), "documento")
        self.assertEqual(sug("linkedin_pdf", "Ana Exemplo\nAnalista de dados"), "curriculo")
        self.assertIsNone(sug("linkedin_pdf", li))
        self.assertIsNone(sug("curriculo_pdf", ""))
        self.assertIsNone(sug("linkedin_zip", "qualquer", "anexos/a.zip"))
        self.assertIsNone(sug("texto", cert, ""))
        m = pp.sugestao_de_area({"tipo": "curriculo_pdf", "texto": cert, "arquivo": "anexos/a.pdf"})
        self.assertEqual(m["motivo"], "Parece um certificado.")
        # sinais vistos em certificados reais: declaração, "concluiu", só a carga horária, o nome do arquivo
        self.assertEqual(sug("curriculo_pdf", "DECLARAÇÃO\nconfere a ANA EXEMPLO declaração de participação no curso"),
                         "documento")
        self.assertEqual(sug("curriculo_pdf", "ANA EXEMPLO\nConcluiu todos os módulos, totalizando 5 horas"), "documento")
        self.assertEqual(sug("curriculo_pdf", "Ana Exemplo\n5/11/19 4/11/20\n32 horas\nCurso Exemplo"), "documento")
        com_nome = lambda nome, texto: (pp.sugestao_de_area(
            {"tipo": "curriculo_pdf", "texto": texto, "arquivo": "anexos/a.pdf", "nome": nome}) or {}).get("area")
        self.assertEqual(com_nome("Escola - Certificado_Curso.pdf", "ANA EXEMPLO\nFUNDAMENTOS\n16 08 2022"), "documento")
        self.assertEqual(com_nome("Faculdade_diploma.pdf", "Ana Exemplo"), "documento")
        self.assertIsNone(com_nome("Curriculo Ana.pdf", "Ana Exemplo\nAnalista de dados"))
        # marca forte de documento acadêmico vale em qualquer tamanho (diploma com histórico tem muitas páginas)
        historico = "HISTÓRICO ESCOLAR\nCurso: Tecnologia\n" + "Disciplina Exemplo 80h aprovado\n" * 300
        m = pp.sugestao_de_area({"tipo": "curriculo_pdf", "texto": historico, "arquivo": "anexos/h.pdf"})
        self.assertEqual((m["area"], m["motivo"]), ("documento", "Parece um diploma ou histórico escolar."))
        self.assertIsNone(sug("curriculo_pdf", cv_longo + "\nhoras extras em projetos de 40 horas"))  # longo: horas não basta

    def test_material_para_pagina(self):
        e = pp.estado_vazio()
        m = pp.adicionar_material(e, "cert.pdf", pdf_minimo(CERTIFICADO))
        p = pp.material_para_pagina(m)
        self.assertEqual((p["area"], p["ocr"], p["sugestao"]["area"]), ("curriculo", False, "documento"))
        pp.mover_material(e, m["id"], "documento")
        p = pp.material_para_pagina(m)
        self.assertEqual((p["tipo"], p["area"], p["sugestao"]), ("documento_pdf", "documento", None))


class TestOcrNoPdf(Base):
    """Spec 037: PDF só de imagem lido pelo OCR do Windows (simulado aqui)."""

    def com_ocr(self, texto, avisos=()):
        ler = mock.patch.object(pp.ocr, "ler_pdf", return_value=(texto, list(avisos)))
        disp = mock.patch.object(pp.ocr, "disponivel", return_value=True)
        return ler, disp

    def test_pdf_sem_texto_lido_por_ocr(self):
        ler, disp = self.com_ocr(f"CERTIFICADO\nCertificamos que Ana Exemplo, CPF {CPF}, concluiu o curso Scrum")
        with ler as chamado, disp:
            r = pp.extrair(self.arquivo("scan.pdf", pdf_minimo([])), area="documento")
        chamado.assert_called_once()
        self.assertTrue(r["ocr"])
        self.assertIn("concluiu o curso Scrum", r["texto"])
        self.assertNotIn(CPF, r["texto"])
        self.assertEqual(r["avisos"], [pp.AVISO_OCR])
        self.assertEqual(r["tipo"], "documento_pdf")

    def test_ocr_vazio_e_pdf_com_texto(self):
        ler, disp = self.com_ocr("", ["OCR do Windows não leu o arquivo (x)"])
        with ler, disp:
            r = pp.extrair(self.arquivo("scan.pdf", pdf_minimo([])))
        self.assertFalse(r["ocr"])
        self.assertTrue(any("imagem" in a for a in r["avisos"]))
        self.assertIn("OCR do Windows não leu o arquivo (x)", r["avisos"])
        ler, disp = self.com_ocr("não devia")
        with ler as chamado, disp:
            r = pp.extrair(self.arquivo("cv.pdf", pdf_minimo(["Ana Exemplo", "Analista"])))
        chamado.assert_not_called()
        self.assertFalse(r["ocr"])

    def test_reler_material(self):
        e = pp.estado_vazio()
        m = pp.adicionar_material(e, "scan.pdf", pdf_minimo([]), area="documento")
        self.assertEqual(m["texto"], "")
        ler, disp = self.com_ocr("Certificado de Conclusão\nCertificamos que Ana Exemplo concluiu o curso")
        with ler, disp:
            pp.reler_material(e, m["id"])
        self.assertTrue(m["ocr"])
        self.assertIn("Certificamos", m["texto"])
        self.assertEqual(m["tipo"], "documento_pdf")  # a área fica
        with self.assertRaisesRegex(pp.PerfilErro, "só para PDF sem texto"):
            pp.reler_material(e, m["id"])  # já tem texto
        cv = pp.adicionar_material(e, "cv.docx", docx_bytes(["Ana"]))
        with self.assertRaisesRegex(pp.PerfilErro, "só para PDF sem texto"):
            pp.reler_material(e, cv["id"])


class TestMateriais(Base):
    def test_guarda_sem_sobrescrever_e_remove(self):
        e = pp.estado_vazio()
        a = pp.adicionar_material(e, "Currículo Ana.docx", docx_bytes(["Ana"]))
        b = pp.adicionar_material(e, "Currículo Ana.docx", docx_bytes(["Ana 2"]))
        self.assertNotEqual(a["arquivo"], b["arquivo"])
        self.assertTrue((self.tmp / a["arquivo"]).exists() and (self.tmp / b["arquivo"]).exists())
        self.assertEqual([m["id"] for m in e["materiais"]], ["m1", "m2"])
        pp.remover_material(e, "m1")
        self.assertEqual([m["id"] for m in e["materiais"]], ["m2"])
        with self.assertRaises(pp.PerfilErro):
            pp.remover_material(e, "m9")

    def test_texto_colado_e_limites(self):
        e = pp.estado_vazio()
        m = pp.adicionar_material(e, "", texto=f"Sou analista. CPF {CPF}")
        self.assertEqual(m["tipo"], "texto")
        self.assertNotIn(CPF, m["texto"])
        with self.assertRaises(pp.PerfilErro):
            pp.adicionar_material(e, "grande.pdf", b"x" * (pp.LIMITE_ARQUIVO + 1))
        with self.assertRaises(pp.PerfilErro):
            pp.adicionar_material(e, "antigo.doc", b"x")
        with self.assertRaises(pp.PerfilErro):
            pp.adicionar_material(e, "quebrado.zip", b"PK nada")
        self.assertEqual(len(e["materiais"]), 1)
        self.assertFalse(list(pp.ANEXOS.glob("*.zip")) if pp.ANEXOS.exists() else [])

    def test_progresso_e_recomecar(self):
        e = pp.carregar()
        pp.adicionar_material(e, "cv.docx", docx_bytes(["Ana"]))
        e["respostas"] = {"cidade": "Curitiba"}
        pp.salvar(e)
        self.assertEqual(pp.carregar()["respostas"], {"cidade": "Curitiba"})
        self.assertEqual(pp.recomecar(), pp.estado_vazio())
        self.assertEqual(pp.carregar()["materiais"], [])
        self.assertTrue(list(pp.ANEXOS.glob("*.docx")))  # os arquivos ficam


class TestRascunho(Base):
    def estado(self):
        e = pp.estado_vazio()
        pp.adicionar_material(e, "cv.docx", docx_bytes(["Analista de Dados — Empresa Alfa (03/2021–atual)",
                                                        f"CPF {CPF}"]))
        pp.adicionar_material(e, "linkedin.zip", zip_linkedin())
        return e

    def test_com_ia_fonte_conflito_e_sem_documentos(self):
        resposta = {
            "objetivo": [{"texto": "Analista de Dados", "fonte": "cv.docx"}],
            "resumo": [{"texto": "Inventado", "fonte": "chute"}, {"texto": "Sem fonte"}],
            "experiencias": [{"cargo": "Analista de Dados", "empresa": "Empresa Alfa", "inicio": "{{conflito:0}}",
                              "fim": "", "itens": [{"texto": "Painéis de vendas", "fonte": "linkedin.zip"}],
                              "fonte": "cv.docx"},
                             {"cargo": "Fantasma", "empresa": "X", "fonte": "inexistente"}],
            "conflitos": [{"campo": "Início na Empresa Alfa", "opcoes": [{"valor": "03/2021", "fonte": "cv.docx"},
                                                                          {"valor": "Mar 2021",
                                                                           "fonte": "linkedin.zip"}]}],
        }
        pedidos = []
        with mock.patch.object(ia, "responder", lambda p, **k: pedidos.append(p) or json.dumps(resposta)):
            r = pp.rascunho_com_ia(self.estado())
        self.assertNotIn(CPF, pedidos[0])
        self.assertIn("### Fonte: cv.docx", pedidos[0])
        self.assertEqual(r["resumo"], [])
        self.assertEqual([x["cargo"] for x in r["experiencias"]], ["Analista de Dados"])
        self.assertEqual(len(r["conflitos"][0]["opcoes"]), 2)
        self.assertIsNone(r["conflitos"][0]["escolha"])

    def test_pedido_com_area_e_ocr(self):
        e = self.estado()
        doc = pp.adicionar_material(e, "cert.pdf", pdf_minimo(CERTIFICADO), area="documento")
        doc["ocr"] = True
        pedido = pp.pedido_rascunho(e)
        self.assertIn("### Fonte: cv.docx\nÁrea: currículo\n", pedido)
        self.assertIn("### Fonte: linkedin.zip\nÁrea: LinkedIn\n", pedido)
        self.assertIn("### Fonte: cert.pdf\nÁrea: certificado ou documento; texto lido por OCR, pode ter erros de leitura\n",
                      pedido)
        self.assertIn("nunca como experiência", pedido)
        self.assertIn("não use número, data ou nome que pareça mal lido", pedido)

    def test_com_ia_erros(self):
        with mock.patch.object(ia, "responder", side_effect=ia.IAErro("sem conexão")):
            with self.assertRaises(pp.PerfilErro):
                pp.rascunho_com_ia(self.estado())
        with mock.patch.object(ia, "responder", return_value="não sei"):
            with self.assertRaises(pp.PerfilErro):
                pp.rascunho_com_ia(self.estado())
        with self.assertRaises(pp.PerfilErro):
            pp.rascunho_com_ia(pp.estado_vazio())

    def test_sem_ia_pelo_zip(self):
        r = pp.rascunho_sem_ia(self.estado())
        self.assertEqual(r["objetivo"][0]["texto"], "Analista de Dados")
        self.assertEqual(r["experiencias"][0]["itens"][1]["texto"], "Automação de relatórios")
        self.assertEqual([x["texto"] for x in r["habilidades"]], ["SQL", "Power BI"])
        self.assertTrue(all(x["fonte"] == "linkedin.zip" for x in r["habilidades"]))


class TestAnamneseEPerfil(Base):
    def rascunho(self):
        return {"objetivo": [], "resumo": [{"texto": "Analista com 5 anos.", "fonte": "cv"}], "formacao": [],
                "certificacoes": [], "habilidades": [], "idiomas": [],
                "experiencias": [{"cargo": "Analista", "empresa": "Alfa", "inicio": "{{conflito:0}}", "fim": "",
                                  "itens": [{"texto": "Painéis", "fonte": "cv"}], "fonte": "cv"},
                                 {"cargo": "Assistente", "empresa": "Beta", "inicio": "2015", "fim": "06/2018",
                                  "itens": [], "fonte": "cv"}],
                "conflitos": [{"campo": "Início", "opcoes": [{"valor": "03/2021", "fonte": "cv"},
                                                             {"valor": "04/2021", "fonte": "zip"}], "escolha": None}]}

    def test_perguntas_por_cargo_e_condicionais(self):
        ids = [p["id"] for p in pp.perguntas_para({"rascunho": self.rascunho()})]
        self.assertIn("conquista:0", ids)
        self.assertIn("medida:1", ids)
        self.assertNotIn("conquista:2", ids)
        cond = {p["id"] for p in pp.PERGUNTAS if p.get("condicao") == pp.EXTERIOR}
        self.assertEqual(cond, {"exterior_remoto", "exterior_paises", "morar_fora", "passaporte", "autorizacao",
                                "sponsor", "fuso", "contratacao"})

    def test_diagnostico(self):
        resp = {"ferramentas": ["Jira — avançado", "Excel"], "conquista:0": "Reduzi o prazo", "medida:0": ""}
        achados = pp.diagnosticar(self.rascunho(), resp)
        textos = " ".join(a["mensagem"] for a in achados)
        self.assertIn("cargos que você busca", textos)
        self.assertIn("Datas sem mês em Assistente — Beta", textos)
        self.assertIn("Sem descrição do que você fazia em Assistente — Beta", textos)
        self.assertIn("Ferramenta sem nível: Excel", textos)
        self.assertNotIn("Jira", textos)
        self.assertIn("sem medida", textos)
        self.assertIn("Conflito sem escolha", textos)
        self.assertIn("Formação vazia", textos)
        self.assertNotIn("Habilidades e ferramentas vazias", textos)  # as ferramentas foram respondidas
        self.assertTrue(all(a["pergunta"] for a in achados))
        ir = {a["mensagem"]: a["ir"] for a in achados}
        self.assertEqual(ir["Ferramenta sem nível: Excel."], "pergunta:ferramentas")
        self.assertEqual(ir["Conflito sem escolha: Início."], "conflitos")
        self.assertEqual(ir["A conquista em Analista está sem medida."], "pergunta:medida:0")

    def test_intervalo(self):
        x = [{"texto": "x", "fonte": "cv"}]
        r = {"experiencias": [{"cargo": "A", "inicio": "01/2015", "fim": "12/2016", "itens": x},
                              {"cargo": "B", "inicio": "Mar 2018", "fim": "", "itens": x}],
             "resumo": x, "formacao": x, "habilidades": x, "idiomas": x}
        achados = pp.diagnosticar(r, {"cargos_alvo": ["X"]})
        self.assertEqual([a["mensagem"] for a in achados], ["Intervalo de 15 meses entre A e B."])

    def test_montar_gravar_e_versao_anterior(self):
        r = self.rascunho()
        resp = {"cargos_alvo": ["Analista de Dados"], "modelos": ["remoto", "hibrido"], "cidade": "Curitiba",
                "estado": "PR", "conquista:0": "Reduzi o prazo do fechamento", "medida:0": "de 10 para 6 dias",
                "interesse_exterior": True, "exterior_remoto": True, "passaporte": False, "nota": f"CPF {CPF}"}
        md = pp.montar_perfil(r, resp, {"email": "ana@exemplo.com"})
        self.assertIn("{{conflito:0}}", md)
        with self.assertRaises(pp.PerfilErro):
            pp.gravar_perfil(md, self.cfg)
        r["conflitos"][0]["escolha"] = 1
        md = pp.montar_perfil(r, resp, {"email": "ana@exemplo.com"})
        self.assertIn("### Analista — Alfa (04/2021–atual)", md)
        self.assertIn("- Maior conquista: Reduzi o prazo do fechamento (medida: de 10 para 6 dias)", md)
        self.assertIn("- E-mail: ana@exemplo.com", md)
        self.assertIn("## Trabalho no exterior", md)
        self.assertIn("- Passaporte válido: não", md)
        self.assertNotIn(CPF, md)
        destino, anterior = pp.gravar_perfil(md, self.cfg)
        self.assertIsNone(anterior)
        self.assertEqual(destino.read_text(encoding="utf-8"), md)
        novo = md.replace("Curitiba", "Londrina")
        destino, anterior = pp.gravar_perfil(novo + f"\nCPF {CPF}\n", self.cfg)
        self.assertEqual(anterior.read_text(encoding="utf-8"), md)
        self.assertNotIn(CPF, destino.read_text(encoding="utf-8"))
        dif = pp.diferencas(md, novo)
        self.assertIn({"tipo": "+", "texto": "- Cidade onde moro: Londrina, PR"}, dif)
        self.assertIn({"tipo": "-", "texto": "- Cidade onde moro: Curitiba, PR"}, dif)

    def test_objetivo_sem_repetir_cargo(self):
        r = {"objetivo": [{"texto": "Analista de Dados", "fonte": "zip"}, {"texto": "Foco em varejo", "fonte": "cv"}]}
        md = pp.montar_perfil(r, {"cargos_alvo": ["analista de dados"]})
        self.assertNotIn("- Analista de Dados", md)
        self.assertIn("- Foco em varejo", md)

    def test_sem_exterior_nao_ha_bloco(self):
        md = pp.montar_perfil({}, {"cargos_alvo": ["X"], "interesse_exterior": False})
        self.assertNotIn("Trabalho no exterior", md)
        self.assertNotIn("Contato", md)


class TestIdiomas(unittest.TestCase):
    """Spec 036: a pergunta de idiomas em linhas de idioma e nível (CEFR ou Nativo), sem estimar nível."""

    def ler(self, texto):
        x = pp.ler_idioma(texto)
        return x["idioma"], x["outro"], x["nivel"], bool(x["original"])

    def test_formato_novo_e_siglas(self):
        self.assertEqual(pp.ler_idioma("Inglês — B2 (Intermediário superior)"),
                         {"idioma": "Inglês", "outro": "", "nivel": "B2", "original": ""})
        self.assertEqual(self.ler("English: C1"), ("Inglês", "", "C1", False))
        self.assertEqual(self.ler("Espanhol - b1"), ("Espanhol", "", "B1", False))
        self.assertEqual(self.ler("Francês — A1 (Iniciante)"), ("Francês", "", "A1", False))
        self.assertEqual(self.ler("Alemão — C2 (Proficiência plena)"), ("Alemão", "", "C2", False))

    def test_nativo(self):
        self.assertEqual(self.ler("Português — nativo"), ("Português (BR)", "", "Nativo", False))
        self.assertEqual(self.ler("Portuguese (Native or bilingual)"), ("Português (BR)", "", "Nativo", False))
        self.assertEqual(self.ler("Português (BR) — Nativo"), ("Português (BR)", "", "Nativo", False))

    def test_nivel_sem_equivalencia_fica_em_branco(self):
        x = pp.ler_idioma("Inglês — Professional Working")
        self.assertEqual((x["idioma"], x["nivel"]), ("Inglês", ""))
        self.assertEqual(x["original"], "Inglês — Professional Working")
        self.assertEqual(self.ler("Inglês fluente"), ("Inglês", "", "", True))
        self.assertEqual(self.ler("Inglês — B2B e SaaS"), ("Inglês", "", "", True))  # sigla só como palavra solta

    def test_outro(self):
        self.assertEqual(self.ler("Italiano — A2 (Básico)"), ("Outro", "Italiano", "A2", False))
        self.assertEqual(self.ler("Português (Portugal) — C2"), ("Outro", "Português (Portugal)", "C2", False))
        self.assertEqual(self.ler("Libras — intermediário"), ("Outro", "Libras", "", True))
        self.assertLessEqual(len(pp.ler_idioma("x" * 100 + " — A1")["outro"]), 60)

    def test_lista(self):
        self.assertEqual(pp.ler_idiomas(None), [])
        self.assertEqual(pp.ler_idiomas([]), [])
        self.assertEqual([x["idioma"] for x in pp.ler_idiomas(["Inglês — B2", "", "Alemão — A2"])], ["Inglês", "Alemão"])

    def test_pergunta(self):
        ps = {p["id"]: p for p in pp.perguntas_para({})}
        p = ps["idiomas"]
        self.assertEqual(p["tipo"], "lista")  # a página antiga continua com a caixa de texto
        self.assertEqual(p["idiomas"], ["Inglês", "Português (BR)", "Francês", "Alemão", "Espanhol"])
        self.assertEqual([n["valor"] for n in p["niveis"]], ["A1", "A2", "B1", "B2", "C1", "C2", "Nativo"])
        rot = {n["valor"]: n["rotulo"] for n in p["niveis"]}
        self.assertEqual((rot["B2"], rot["Nativo"]), ("B2 (Intermediário superior)", "Nativo"))
        self.assertTrue(all(n["descricao"] for n in p["niveis"]))
        self.assertEqual(p["linhas"], [])
        p = {x["id"]: x for x in pp.perguntas_para({"respostas": {"idiomas": ["Inglês — B2", "Libras"]}})}["idiomas"]
        self.assertEqual([(x["idioma"], x["nivel"]) for x in p["linhas"]], [("Inglês", "B2"), ("Outro", "")])
        self.assertNotIn("linhas", next(x for x in pp.PERGUNTAS if x["id"] == "idiomas"))  # o catálogo não muda

    def test_linhas_do_rascunho_sem_resposta(self):
        rascunho = {"idiomas": [{"texto": "Português — nativo", "fonte": "zip"},
                                {"texto": "Inglês — Professional Working", "fonte": "zip"}]}
        p = {x["id"]: x for x in pp.perguntas_para({"rascunho": rascunho})}["idiomas"]
        self.assertEqual([(x["idioma"], x["nivel"], bool(x["original"])) for x in p["linhas"]],
                         [("Português (BR)", "Nativo", False), ("Inglês", "", True)])
        p = {x["id"]: x for x in pp.perguntas_para({"rascunho": rascunho, "respostas": {"idiomas": ["Alemão — A2"]}})}
        self.assertEqual([x["idioma"] for x in p["idiomas"]["linhas"]], ["Alemão"])  # a resposta vale

    def test_resposta_substitui_o_idioma_do_rascunho(self):
        rascunho = {"idiomas": [{"texto": "Inglês — Professional Working", "fonte": "zip"},
                                {"texto": "Alemão — Elementary", "fonte": "zip"}]}
        md = pp.montar_perfil(rascunho, {"idiomas": ["Inglês — B2 (Intermediário superior)"]})
        self.assertIn("- Inglês — B2 (Intermediário superior)", md)
        self.assertNotIn("Professional Working", md)
        self.assertIn("- Alemão — Elementary", md)  # idioma fora da resposta continua

    def test_sem_pergunta_de_ingles(self):
        self.assertNotIn("ingles", [p["id"] for p in pp.perguntas_para({})])
        self.assertEqual(pp.validar_respostas({"ingles": "avancado"}), {"ingles": "avancado"})  # id antigo aceito

    def test_ingles_uma_vez_no_perfil(self):
        md = pp.montar_perfil({}, {"idiomas": ["Inglês — B2 (Intermediário superior)"], "ingles": "avancado"})
        self.assertIn("- Inglês — B2 (Intermediário superior)", md)
        self.assertNotIn("nível declarado", md)
        md = pp.montar_perfil({}, {"idiomas": ["Inglês — Professional Working"], "ingles": "avancado"})
        self.assertIn("- Inglês (nível declarado): avancado", md)
        md = pp.montar_perfil({}, {"ingles": "avancado"})
        self.assertIn("- Inglês (nível declarado): avancado", md)


class TestFiltros(unittest.TestCase):
    def test_propoe_a_partir_das_respostas(self):
        resp = {"cargos_alvo": ["Analista de Dados", "BI"], "cargos_alvo_en": ["Data Analyst"],
                "modelos": ["remoto", "hibrido"], "cidade": "Curitiba", "estado": "PR", "senioridade": ["pleno"],
                "interesse_exterior": True, "exterior_remoto": True, "exterior_paises": ["Portugal", "Marte"]}
        r = pp.propor_filtros(resp, {})
        f = r["filtros"]
        self.assertEqual(f["termos"], ['"Analista de Dados"', "BI", '"Data Analyst"'])
        self.assertEqual(f["localidade"]["cidade"], "Curitiba")
        self.assertEqual(f["modelos"], {"remoto": True, "hibrido": True, "presencial": False})
        self.assertEqual({k: f["internacional"][k] for k in ("ativo", "paises", "termos")},
                         {"ativo": True, "paises": ["Portugal"], "termos": ['"Data Analyst"']})
        self.assertEqual(f["senioridades"], ["pleno"])
        self.assertIn("Cargos", [m["campo"] for m in r["mudancas"]])

    def test_sem_cidade_fica_remoto_e_sem_exterior_desligado(self):
        r = pp.propor_filtros({"cargos_alvo": ["Analista"], "modelos": ["hibrido"], "interesse_exterior": True}, {})
        self.assertEqual(r["filtros"]["modelos"], {"remoto": True, "hibrido": False, "presencial": False})
        self.assertFalse(r["filtros"]["internacional"]["ativo"])
        self.assertEqual(len(r["avisos"]), 2)

    def test_sem_cargos(self):
        with self.assertRaises(pp.PerfilErro):
            pp.propor_filtros({}, {})


class TestProgressoEPrevia(Base):
    def estado(self):
        e = pp.estado_vazio()
        pp.adicionar_material(e, "cv.docx", docx_bytes(["Analista de Dados na Alfa", f"CPF {CPF}"]))
        e["rascunho"] = {"origem": "sem_ia", "objetivo": [], "resumo": [], "formacao": [], "certificacoes": [],
                         "habilidades": [], "idiomas": [],
                         "experiencias": [{"cargo": "Analista", "empresa": "Alfa", "inicio": "{{conflito:0}}", "fim": "",
                                           "itens": [], "fonte": "cv.docx"}],
                         "conflitos": [{"campo": "Início", "opcoes": [{"valor": "03/2021", "fonte": "cv.docx"}],
                                        "escolha": None}]}
        return e

    def test_atualizar_progresso(self):
        e = self.estado()
        pp.atualizar_progresso(e, {"etapa": "anamnese", "respostas": {"cidade": f"Curitiba CPF {CPF}", "passaporte": True},
                                   "escolhas": {"0": 0}, "contato": {"email": "ana@exemplo.com", "linkedin": ""}})
        self.assertEqual(e["etapa"], "anamnese")
        self.assertNotIn(CPF, e["respostas"]["cidade"])
        self.assertIs(e["respostas"]["passaporte"], True)
        self.assertEqual(e["rascunho"]["conflitos"][0]["escolha"], 0)
        self.assertEqual(e["contato"], {"email": "ana@exemplo.com"})
        pp.atualizar_progresso(e, {"respostas": {"cidade": None}})
        self.assertIsNone(e["respostas"]["cidade"])
        self.assertIs(e["respostas"]["passaporte"], True)
        for ruim in ({"etapa": "outra"}, {"respostas": {"inventada": "x"}}, {"escolhas": {"0": 5}}, {"escolhas": {"3": 0}},
                     {"escolhas": {"0": True}}, {"contato": {"email": "nao-e-email"}},
                     {"contato": {"linkedin": "https://exemplo.com/ana"}}, {"contato": {"telefone": "abc"}}):
            with self.assertRaises(pp.PerfilErro, msg=ruim):
                pp.atualizar_progresso(e, ruim)

    def test_previa_novo_e_existente(self):
        e = self.estado()
        p = pp.previa(e, cfg=self.cfg)
        self.assertIsNone(p["diferencas"])
        self.assertEqual(p["conflitos_pendentes"], 1)
        self.assertEqual(p["textos"][0]["nome"], "cv.docx")
        self.assertNotIn(CPF, json.dumps(p, ensure_ascii=False))
        (self.tmp / "perfil.md").write_text("# Meu perfil de carreira\n\n## Objetivo\n- Cargos que procuro: Antigo\n",
                                            encoding="utf-8")
        p = pp.previa(e, cfg=self.cfg)  # sem IA e com perfil: a revisão começa do perfil atual
        self.assertIn("Cargos que procuro: Antigo", p["markdown"])
        self.assertEqual(p["textos"][0]["nome"], "Montado a partir das suas respostas")
        self.assertEqual(p["diferencas"], [])
        p = pp.previa(e, markdown="# Meu perfil de carreira\n\n## Objetivo\n- Cargos que procuro: Novo\n", cfg=self.cfg)
        self.assertIn({"tipo": "+", "texto": "- Cargos que procuro: Novo"}, p["diferencas"])
        e["rascunho"]["origem"] = "ia"
        self.assertNotIn("Antigo", pp.previa(e, cfg=self.cfg)["markdown"])  # com IA, o rascunho já partiu do perfil atual

    def test_sugestao_em_ingles(self):
        resp = {"cargos_alvo": ["Analista de Dados"], "interesse_exterior": True, "exterior_remoto": True,
                "exterior_paises": ["Portugal"]}
        with mock.patch.object(ia, "responder", return_value='["Data Analyst"]'):
            r = pp.propor_filtros(resp, {}, sugerir_en=True)
        self.assertTrue(r["sugestao_en"])
        self.assertEqual(r["filtros"]["internacional"]["termos"], ['"Data Analyst"'])
        self.assertTrue(any("sugestão da IA" in a for a in r["avisos"]))
        with mock.patch.object(ia, "responder", side_effect=ia.IAErro("fora do ar")):
            with self.assertRaises(pp.PerfilErro):
                pp.propor_filtros(resp, {}, sugerir_en=True)


class TestLinhaDeComando(Base):
    def test_extrair_e_diagnostico(self):
        cv = self.arquivo("cv.docx", docx_bytes(["Ana", f"CPF {CPF}"]))
        saida = io.StringIO()
        with mock.patch("sys.stdout", saida):
            self.assertEqual(pp.main(["extrair", str(cv)]), 0)
        self.assertNotIn(CPF, saida.getvalue())
        self.assertEqual(json.loads(saida.getvalue())["tipo"], "curriculo_docx")
        saida = io.StringIO()
        with mock.patch("sys.stdout", saida):
            self.assertEqual(pp.main(["extrair", str(cv), "--area", "documento"]), 0)
        self.assertEqual(json.loads(saida.getvalue())["tipo"], "documento_docx")
        perfil = self.tmp / "perfil.md"
        perfil.write_text("# Meu perfil\n\n## Objetivo\n- Cargos que procuro: Analista\n\n## Experiência\n\n"
                          "### Analista — Alfa (2019–atual)\n", encoding="utf-8")
        saida = io.StringIO()
        with mock.patch("sys.stdout", saida):
            self.assertEqual(pp.main(["diagnostico", str(perfil)]), 0)
        self.assertIn("Datas sem mês em Analista — Alfa", saida.getvalue())
        self.assertNotIn("cargos que você busca", saida.getvalue())
        self.assertIn("Formação vazia", saida.getvalue())
        completo = self.tmp / "completo.md"
        completo.write_text("# Meu perfil\n\n## Objetivo\n- Cargos que procuro: Analista\n\n## Resumo\nAnalista.\n\n"
                            "## Experiência\n\n### Analista — Alfa (03/2019–atual)\n- Painéis\n\n## Formação\n- Estatística\n\n"
                            "## Habilidades e ferramentas (com nível)\n- SQL — avançado\n\n## Idiomas\n- Inglês — leitura\n",
                            encoding="utf-8")
        saida = io.StringIO()
        with mock.patch("sys.stdout", saida):
            self.assertEqual(pp.main(["diagnostico", str(completo)]), 0)
        self.assertIn("Nada a apontar.", saida.getvalue())
        with mock.patch("sys.stdout", io.StringIO()), mock.patch.object(pp, "_config", lambda: self.cfg):
            self.assertEqual(pp.main(["gravar", str(completo)]), 0)
            self.assertEqual(pp.main(["gravar", str(completo)]), 0)
        self.assertEqual((self.tmp / "perfil.md").read_text(encoding="utf-8"), completo.read_text(encoding="utf-8"))
        self.assertEqual(len(list(pp.ANTERIORES.glob("perfil-*.md"))), 2)  # o perfil.md do teste acima e o primeiro gravado
        with mock.patch("sys.stderr", io.StringIO()):
            self.assertEqual(pp.main(["extrair", str(self.tmp / "nao-existe.pdf")]), 1)
            pendente = self.arquivo("pendente.md", "### A — B ({{conflito:0}}–atual)\n".encode())
            self.assertEqual(pp.main(["gravar", str(pendente)]), 1)


if __name__ == "__main__":
    unittest.main()
