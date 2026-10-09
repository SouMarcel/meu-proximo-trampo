#!/usr/bin/env python3
"""Conferência do currículo e da carta antes de entregar, sem IA: fatos, requisitos da vaga e ATS.

  python conferir.py curriculos/base.json                      fatos contra o perfil e ATS do .docx
  python conferir.py curriculos/x.json --vaga-id ID            também os requisitos de uma vaga do dashboard
  python conferir.py curriculos/x.json --vaga vaga.txt         idem, com o texto da vaga num arquivo
  python conferir.py curriculos/x.json --perfil outro.md --json
  python conferir.py --carta carta.txt --vaga-id ID [--respostas respostas.json]   carta de apresentação

Veredito: ok, conferir ou bloquear (o mais grave dos pontos). Bloqueia número que não está no perfil,
lacuna da vaga listada como competência e dado pessoal no formato dos Estados Unidos. Códigos de saída:
0 ok, 1 conferir, 3 bloquear, 2 erro de uso. O .docx conferido é o que está ao lado do JSON.

Na carta: número ou data que não estão no perfil nem nas respostas da pessoa bloqueiam; tamanho fora de
350 a 420 palavras, empresa da vaga não citada, nenhum requisito do anúncio citado, expressões vazias e
nome próprio que não aparece no perfil nem na vaga ficam para conferir.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

RAIZ = Path(__file__).resolve().parent
NIVEIS = {"ok": 0, "conferir": 1, "bloquear": 2}
FONTES_COMUNS = {"calibri", "arial", "helvetica", "times new roman", "georgia", "garamond", "cambria", "verdana",
                 "aptos", "segoe ui", "tahoma", "book antiqua", "palatino linotype", "liberation sans", "carlito"}
MESES = {"jan": 1, "fev": 2, "feb": 2, "mar": 3, "abr": 4, "apr": 4, "mai": 5, "may": 5, "jun": 6, "jul": 7,
         "ago": 8, "aug": 8, "set": 9, "sep": 9, "out": 10, "oct": 10, "nov": 11, "dez": 12, "dec": 12}
VAZIAS = set("""a o as os um uma uns umas de do da dos das em no na nos nas por para com sem e ou que se ao aos à às
como mais menos muito boa bom bons boas forte fortes solido solida experiencia experiencias conhecimento conhecimentos
vivencia vivencias pratica dominio familiaridade habilidade habilidades capacidade nivel avancado avancada intermediario
basico desejavel desejaveis diferencial diferenciais obrigatorio requisito requisitos ano anos mes meses minimo area
the an of in on at to for with without and or as is are be have has strong solid good great excellent knowledge
experience experienced proven ability skills skill familiarity understanding hands years year plus preferred required
nice must using use etc e/ou tipo tipos atuacao atuar trabalhar trabalho saber ter possuir""".split())
CABECALHOS_REQ = ("requisitos", "qualificacoes", "o que buscamos", "o que esperamos", "o que voce precisa",
                  "o que procuramos", "para se candidatar", "desejavel", "desejaveis", "diferenciais", "competencias",
                  "requirements", "qualifications", "must have", "nice to have", "what you'll need", "what you need",
                  "what we're looking for", "what we are looking for", "you have", "you bring", "skills")
CABECALHOS_OUTROS = ("responsabilidades", "atividades", "atribuicoes", "beneficios", "sobre", "quem somos",
                     "o que oferecemos", "local", "jornada", "responsibilities", "what you'll do", "benefits",
                     "about", "perks", "we offer", "salario", "remuneracao")
FAMILIAS = [{"jira", "azure boards", "azure devops", "trello", "clickup", "asana"},
            {"power bi", "tableau", "looker", "metabase", "qlik", "data studio"},
            {"scrum", "kanban", "agil", "agile", "safe", "lean"},
            {"sql", "postgresql", "mysql", "sql server", "oracle"},
            {"aws", "azure", "gcp", "google cloud"},
            {"figma", "sketch", "adobe xd"},
            {"excel", "google sheets", "planilhas", "spreadsheets"}]
NUMEROS_PALAVRAS = {"um": 1, "uma": 1, "dois": 2, "duas": 2, "tres": 3, "quatro": 4, "cinco": 5, "seis": 6, "sete": 7,
                    "oito": 8, "nove": 9, "dez": 10, "onze": 11, "doze": 12, "treze": 13, "quatorze": 14, "catorze": 14,
                    "quinze": 15, "dezesseis": 16, "dezessete": 17, "dezoito": 18, "dezenove": 19, "vinte": 20,
                    "trinta": 30, "quarenta": 40, "cinquenta": 50, "sessenta": 60, "setenta": 70, "oitenta": 80,
                    "noventa": 90, "cem": 100, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
                    "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12, "fifteen": 15,
                    "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "sixty": 60, "seventy": 70, "eighty": 80,
                    "ninety": 90, "hundred": 100}
DADOS_PESSOAIS = re.compile(r"\b(nascimento|nascid[oa]|born|date of birth|casad[oa]|solteir[oa]|married|single|"
                            r"estado civil|marital|nacionalidade|nationality|cpf|rg|ssn|idade|age|anos de idade)\b", re.I)


def norm(texto) -> str:
    """Minúsculas sem acento; outros sinais fora do ASCII (travessão, aspas curvas) viram espaço."""
    t = "".join("" if unicodedata.combining(c) else c if ord(c) < 128 else " "
                for c in unicodedata.normalize("NFKD", str(texto or "")))
    return re.sub(r"\s+", " ", t.lower()).strip()


def termos(texto: str) -> list[str]:
    """Palavras significativas (minúsculas, sem acento), na ordem, sem repetir."""
    vistos = []
    for t in re.findall(r"[a-z0-9][a-z0-9+#.]*[a-z0-9+#]|[a-z0-9]", norm(texto)):
        if t not in VAZIAS and not t.isdigit() and t not in vistos:
            vistos.append(t)
    return vistos


def _contem(base_norm: str, termo: str) -> bool:
    return re.search(r"(?<![a-z0-9])" + re.escape(termo) + r"(?![a-z0-9])", base_norm) is not None


# ---------------------------------------------------------------- texto do currículo

def _lista(v) -> list:
    return [] if v is None else (v if isinstance(v, list) else [v])


def trechos(cv: dict) -> list[tuple[str, str]]:
    """(seção, texto) de tudo que o currículo afirma, sem o contato."""
    out = [("titulo", cv.get("titulo") or "")]
    out += [("resumo", str(p)) for p in _lista(cv.get("resumo"))]
    out += [("destaques", str(p)) for p in _lista(cv.get("destaques"))]
    for g in _lista(cv.get("competencias")):
        if isinstance(g, dict):
            out += [("competencias", str(i)) for i in _lista(g.get("itens"))]
        else:
            out.append(("competencias", str(g)))
    for emp in _lista(cv.get("experiencia")):
        out.append(("experiencia", str(emp.get("descricao") or "")))
        for c in emp.get("cargos") or [{"cargo": emp.get("cargo"), "itens": emp.get("itens")}]:
            out += [("experiencia", str(i)) for i in _lista(c.get("itens"))]
    for pj in _lista(cv.get("projetos")):
        if isinstance(pj, dict):
            out += [("projetos", str(pj.get("descricao") or ""))] + [("projetos", str(i)) for i in _lista(pj.get("itens"))]
        else:
            out.append(("projetos", str(pj)))
    for chave in ("certificacoes", "cursos", "idiomas"):
        out += [(chave, str(i)) for i in _lista(cv.get(chave))]
    for bloco in _lista(cv.get("extras")):
        out += [("extras", str(i)) for i in _lista(bloco.get("itens") if isinstance(bloco, dict) else bloco)]
    return [(s, t) for s, t in out if t.strip()]


# ---------------------------------------------------------------- números e datas

DATA = re.compile(r"\b(\d{1,2})\s*/\s*(\d{4})\b|\b(\d{4})\s*-\s*(\d{1,2})\b|\b([a-zç]{3})[a-zç]*\.?\s*(?:de\s+)?/?\s*(\d{4})\b", re.I)
NUMERO = re.compile(
    r"(?P<moeda>R\$|US\$|U\$|\$|€|£)?\s?"
    r"(?P<n>\d{1,3}(?:[.\s]\d{3})+(?:,\d+)?|\d+(?:[.,]\d+)?)"
    r"(?:\s?(?P<suf>%|por cento|percent|mil\b|k\b|mi\b|milh[oõ]es|milh[aã]o|mm\b|m\b|bi\b|bilh[oõ]es|bilh[aã]o|x\b|vezes|times))?",
    re.I)
MULT = {"mil": 1e3, "k": 1e3, "mi": 1e6, "milhoes": 1e6, "milhao": 1e6, "mm": 1e6, "m": 1e6, "bi": 1e9,
        "bilhoes": 1e9, "bilhao": 1e9}


def _valor(n: str) -> float:
    n = n.replace(" ", "")
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", n):  # 1.500 / 1.500,50
        return float(n.replace(".", "").replace(",", "."))
    if re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", n):  # 1,500 (inglês)
        return float(n.replace(",", ""))
    return float(n.replace(",", "."))


def numeros(texto: str) -> list[tuple[float, str, str]]:
    """(valor, tipo, como está escrito) de cada número do texto, sem datas e anos. Tipos: pct, x, moeda, n."""
    sem_datas = DATA.sub(" ", str(texto or ""))
    sem_datas = re.sub(r"\b(19|20)\d{2}\b", " ", sem_datas)
    achados = []
    for m in NUMERO.finditer(sem_datas):
        try:
            valor = _valor(m.group("n"))
        except ValueError:
            continue
        suf = norm(m.group("suf") or "")
        tipo = "pct" if suf in ("%", "por cento", "percent") else "x" if suf in ("x", "vezes", "times") else \
            "moeda" if m.group("moeda") else "n"
        valor *= MULT.get(suf, 1)
        achados.append((round(valor, 4), tipo, m.group(0).strip()))
    # números por extenso ("trinta por cento", "five times")
    palavras = norm(texto).split()
    i = 0
    while i < len(palavras):
        p = re.sub(r"[^a-z]", "", palavras[i])
        if p in NUMEROS_PALAVRAS and p not in ("um", "uma", "one"):
            valor, j = NUMEROS_PALAVRAS[p], i + 1
            while j + 1 < len(palavras) and palavras[j] in ("e", "and") and re.sub(r"[^a-z]", "", palavras[j + 1]) in NUMEROS_PALAVRAS:
                valor += NUMEROS_PALAVRAS[re.sub(r"[^a-z]", "", palavras[j + 1])]
                j += 2
            resto = " ".join(palavras[j:j + 2])
            tipo = "pct" if resto.startswith(("por cento", "percent")) else "x" if resto.startswith(("vezes", "times")) else "n"
            if tipo != "n":
                achados.append((float(valor), tipo, " ".join(palavras[i:j + (2 if tipo == "pct" and resto.startswith("por cento") else 1)])))
            i = j
        else:
            i += 1
    return achados


def _bate(num: tuple, base: list[tuple]) -> bool:
    valor, tipo, _ = num
    for v, t, _ in base:
        if abs(v - valor) < 1e-6 and (t == tipo or "n" in (t, tipo) or {t, tipo} == {"moeda", "n"}):
            if tipo == "pct" and t not in ("pct", "n"):
                continue
            return True
    return False


def datas(texto: str) -> list[tuple[int, int | None]]:
    """(ano, mês) de cada data do texto; mês None quando só há o ano."""
    out, t, aceitas = [], norm(texto), []
    for m in DATA.finditer(t):
        if m.group(1):
            mes, ano = int(m.group(1)), int(m.group(2))
        elif m.group(3):
            ano, mes = int(m.group(3)), int(m.group(4))
        else:
            if m.group(5)[:3] not in MESES:
                continue  # "since 2015": não é mês, o ano conta abaixo
            mes, ano = MESES[m.group(5)[:3]], int(m.group(6))
        if 1 <= mes <= 12:
            out.append((ano, mes))
            aceitas.append(m.span())
    sem = t
    for a, b in reversed(aceitas):  # só as datas reconhecidas saem antes de procurar os anos soltos
        sem = sem[:a] + " " + sem[b:]
    out += [(int(a), None) for a in re.findall(r"\b((?:19|20)\d{2})\b", sem)]
    return out


# ---------------------------------------------------------------- partes da conferência

def conferir_fatos(cv: dict, perfil: str) -> list[dict]:
    pontos, base_num, perfil_n = [], numeros(perfil), norm(perfil)
    vistos = set()
    for secao, texto in trechos(cv):
        for num in numeros(texto):
            chave = (num[0], num[1])
            if chave in vistos or _bate(num, base_num):
                continue
            vistos.add(chave)
            pontos.append({"tipo": "numero", "valor": num[2], "trecho": texto[:160], "nivel": "bloquear",
                           "mensagem": f"o número {num[2]} não está no perfil: confirme com a pessoa ou tire"})
    limpar = lambda s: re.sub(r"\b(s\.?a\.?|ltda\.?|me|eireli|inc\.?|llc|ltd\.?)\b", "", norm(s)).strip(" .,-")
    for emp in _lista(cv.get("experiencia")):
        nome = limpar(emp.get("empresa"))
        if nome and nome not in perfil_n:
            pontos.append({"tipo": "empresa", "valor": emp.get("empresa"), "trecho": emp.get("empresa"), "nivel": "conferir",
                           "mensagem": "empresa que não aparece no perfil"})
        cargos = emp.get("cargos") or [{"cargo": emp.get("cargo"), "periodo": emp.get("periodo")}]
        for c in cargos:
            cargo = norm(c.get("cargo"))
            if cargo and cargo not in perfil_n:
                pontos.append({"tipo": "cargo", "valor": c.get("cargo"), "trecho": c.get("cargo"), "nivel": "conferir",
                               "mensagem": "cargo que não aparece no perfil (se foi traduzido, confira a tradução)"})
        pontos += _conferir_datas(" ".join(str(x) for x in [emp.get("periodo")] + [c.get("periodo") for c in cargos] if x), perfil)
    return pontos


def _conferir_datas(texto: str, perfil: str) -> list[dict]:
    base = datas(perfil)
    anos, meses = {a for a, _ in base}, {(a, m) for a, m in base if m}
    pontos = []
    for ano, mes in datas(texto):
        if mes and (ano, mes) in meses or not mes and ano in anos:
            continue
        msg = ("o perfil tem só o ano; o mês não está confirmado" if mes and ano in anos
               else "data que não está no perfil")
        pontos.append({"tipo": "data", "valor": f"{mes:02d}/{ano}" if mes else str(ano), "trecho": texto, "nivel": "conferir",
                       "mensagem": msg})
    return pontos


def requisitos_da_vaga(vaga: str) -> list[str]:
    linhas = [l.strip() for l in str(vaga or "").splitlines()]
    reqs, dentro = [], False
    for l in linhas:
        if not l:
            continue
        cab = norm(l).strip("#*:- ")
        curta = len(l) <= 60 and (l.rstrip().endswith(":") or l.startswith("#") or l.isupper() or cab in CABECALHOS_REQ)
        if curta and any(cab.startswith(c) for c in CABECALHOS_REQ):
            dentro = True
            continue
        if curta and any(cab.startswith(c) for c in CABECALHOS_OUTROS):
            dentro = False
            continue
        if dentro:
            reqs.append(re.sub(r"^[\-*•·▪◦‣–\d.)\s]+", "", l).strip())
    if not reqs:  # sem cabeçalho: as linhas com marcador
        reqs = [re.sub(r"^[\-*•·▪◦‣–\s]+", "", l).strip() for l in linhas if re.match(r"^[\-*•·▪◦‣–]\s+", l)]
    return [r for r in reqs if termos(r)][:40]


def _familia(req_n: str, perfil_n: str) -> str | None:
    for fam in FAMILIAS:
        pedidos = [x for x in fam if _contem(req_n, x)]
        if pedidos and not any(_contem(perfil_n, x) for x in pedidos):
            tem = [x for x in fam if _contem(perfil_n, x)]
            if tem:
                return tem[0]
    return None


def conferir_requisitos(cv: dict, perfil: str, vaga: str) -> tuple[list[dict], dict, list[dict]]:
    """(requisitos com situação, cobertura das palavras-chave, pontos de bloqueio por lacuna em Competências)."""
    perfil_n, cv_n = norm(perfil), norm(" ".join(t for _, t in trechos(cv)))
    comp_n = norm(" ".join(t for s, t in trechos(cv) if s == "competencias"))
    linhas_perfil = [l for l in str(perfil).splitlines() if l.strip()]
    saida, pontos, chaves = [], [], []
    for req in requisitos_da_vaga(vaga):
        ts = termos(req)
        chaves += [t for t in ts if t not in chaves]
        achados = [t for t in ts if _contem(perfil_n, t)]
        ratio = len(achados) / len(ts)
        equivalente = _familia(norm(req), perfil_n)
        if ratio >= 0.6:
            situacao = "tem"
        elif achados or equivalente:
            situacao = "sustentado"
        else:
            situacao = "lacuna"
        evid = ""
        if achados:
            evid = max(linhas_perfil, key=lambda l: sum(_contem(norm(l), t) for t in achados)).strip()[:160]
        elif equivalente:
            evid = f"equivalente no perfil: {equivalente}"
        saida.append({"texto": req[:200], "situacao": situacao, "evidencia": evid})
        if situacao == "lacuna":
            na_lista = [t for t in ts if len(t) >= 3 and _contem(comp_n, t)]
            if na_lista:
                pontos.append({"tipo": "lacuna", "valor": ", ".join(na_lista), "trecho": req[:160], "nivel": "bloquear",
                               "mensagem": "requisito que é lacuna aparece em Competências: tire do currículo"})
    achadas = [t for t in chaves if _contem(cv_n, t)]
    cobertura = {"achadas": len(achadas), "total": len(chaves), "faltando": [t for t in chaves if t not in achadas][:15]}
    return saida, cobertura, pontos


def conferir_ats(cv: dict, docx: Path | None) -> list[dict]:
    pontos = []
    formato = (cv.get("tecnicas") or {}).get("formato")
    contato = [str(c) for c in _lista(cv.get("contato"))]
    if not contato:
        pontos.append({"nivel": "conferir", "mensagem": "sem contato no currículo", "trecho": ""})
    if formato == "us":
        for c in contato:
            if DADOS_PESSOAIS.search(c) or re.search(r"\d{3}\.?\d{3}\.?\d{3}-?\d{2}", c):
                pontos.append({"nivel": "bloquear", "trecho": c,
                               "mensagem": "dado pessoal no formato dos Estados Unidos (idade, estado civil, documentos)"})
    for bloco in _lista(cv.get("extras")):
        titulo = str(bloco.get("titulo") or "") if isinstance(bloco, dict) else ""
        if titulo and len(titulo) > 30:
            pontos.append({"nivel": "conferir", "trecho": titulo, "mensagem": "título de seção longo: ATS entende melhor títulos padrão"})
    if not docx or not Path(docx).exists():
        pontos.append({"nivel": "conferir", "trecho": "", "mensagem": "documento .docx não encontrado: ATS do arquivo não conferido"})
        return pontos
    try:
        from docx import Document
    except ImportError:
        pontos.append({"nivel": "conferir", "trecho": "", "mensagem": "falta python-docx: ATS do arquivo não conferido"})
        return pontos
    doc = Document(str(docx))
    corpo = doc.element.body.xml
    if doc.tables:
        pontos.append({"nivel": "conferir", "trecho": "", "mensagem": f"{len(doc.tables)} tabela(s): ATS costuma embaralhar o texto"})
    if "<w:drawing" in corpo or "<w:pict" in corpo or "<pic:" in corpo:
        pontos.append({"nivel": "conferir", "trecho": "", "mensagem": "imagem no documento: ATS não lê"})
    if "w:txbxContent" in corpo:
        pontos.append({"nivel": "conferir", "trecho": "", "mensagem": "caixa de texto: ATS costuma pular o conteúdo"})
    if re.search(r'<w:cols[^>]*w:num="([2-9])"', corpo):
        pontos.append({"nivel": "conferir", "trecho": "", "mensagem": "mais de uma coluna: ATS lê fora de ordem"})
    for sec in doc.sections:
        for parte, nome in ((sec.header, "cabeçalho"), (sec.footer, "rodapé")):
            texto = " ".join(p.text for p in parte.paragraphs if p.text.strip())
            if texto:
                pontos.append({"nivel": "conferir", "trecho": texto[:120],
                               "mensagem": f"texto no {nome}: contato e dados devem ficar no corpo do documento"})
    usados = set(re.findall(r'<w:pStyle w:val="([^"]+)"', corpo)) | {"Normal"}
    estilos = "".join(st.element.xml for st in doc.styles if getattr(st, "style_id", None) in usados or st.name in usados)
    fontes = set(re.findall(r'w:ascii="([^"]+)"', corpo + estilos))  # só as fontes usadas de fato
    for f in sorted(fontes):
        if norm(f) not in FONTES_COMUNS and not f.startswith("+"):
            pontos.append({"nivel": "conferir", "trecho": f, "mensagem": "fonte pouco comum: use Calibri, Arial ou outra padrão"})
    return pontos


# ---------------------------------------------------------------- tudo junto

def conferir(cv: dict, perfil: str, vaga: str | None = None, docx: Path | None = None) -> dict:
    fatos = conferir_fatos(cv, perfil or "")
    requisitos, cobertura, lacunas = conferir_requisitos(cv, perfil or "", vaga) if vaga else ([], None, [])
    ats = conferir_ats(cv, docx)
    todos = fatos + lacunas + ats
    pior = max((NIVEIS[p["nivel"]] for p in todos), default=0)
    veredito = {0: "ok", 1: "conferir", 2: "bloquear"}[pior]
    return {"veredito": veredito, "fatos": fatos + lacunas, "requisitos": requisitos, "cobertura": cobertura, "ats": ats}


# ---------------------------------------------------------------- carta e respostas

PALAVRAS_CARTA = (350, 420)
VAZIAS_EN = set("""about the this that with you your our we will but who how for and are is at least in of to be or
from have has can all any not such as on by an it its their they them per also more than""".split())
EXPRESSOES_VAZIAS = (
    "sou apaixonado", "sou apaixonada", "apaixonado por", "apaixonada por", "proativo", "proativa", "fora da caixa",
    "sinergia", "vestir a camisa", "perfil dinamico", "busco novos desafios", "team player", "self-starter",
    "passionate about", "think outside the box", "fast-paced environment", "go-getter", "hard worker",
    "results-driven", "dynamic environment", "i believe i would be a great fit", "perfect fit",
)
NOMES_COMUNS = {
    "i", "i'm", "i've", "i'd", "i'll", "dear", "hiring", "manager", "team", "sincerely", "regards", "best", "thank",
    "thanks", "hello", "hi", "prezado", "prezada", "prezados", "atenciosamente", "obrigado", "obrigada", "ola",
    "english", "portuguese", "spanish", "ingles", "portugues", "espanhol", "brazil", "brasil", "linkedin",
    "january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november",
    "december", "janeiro", "fevereiro", "marco", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro",
    "novembro", "dezembro", "monday", "friday", "my", "as", "in", "at", "the", "this", "that", "with", "and",
}


def contar_palavras(texto: str) -> int:
    return len(re.findall(r"[^\W_]+(?:[-'’][^\W_]+)*", str(texto or "")))


def _nomes_proprios(texto: str) -> list[str]:
    """Sequências com inicial maiúscula que não começam frase (empresas, cargos, lugares, pessoas)."""
    achados = []
    for frase in re.split(r"(?<=[.!?:;])\s+|\n+", str(texto or "")):
        palavras = frase.split()
        atual = []
        for i, p in enumerate(palavras):
            limpa = p.strip(",.;:!?()\"'“”‘’")
            if i > 0 and re.match(r"^[A-ZÀ-Ý][\w&.-]*$", limpa) and norm(limpa) not in NOMES_COMUNS:
                atual.append(limpa)
            else:
                if atual:
                    achados.append(" ".join(atual))
                atual = []
        if atual:
            achados.append(" ".join(atual))
    return list(dict.fromkeys(achados))


def _fatos_do_texto(texto: str, base: str, nivel_data: str) -> list[dict]:
    """Números e datas do texto que não estão na base (perfil + o que a pessoa respondeu)."""
    pontos, base_num, vistos = [], numeros(base), set()
    for num in numeros(texto):
        if (num[0], num[1]) in vistos or _bate(num, base_num):
            continue
        vistos.add((num[0], num[1]))
        pontos.append({"tipo": "numero", "valor": num[2], "nivel": "bloquear",
                       "mensagem": f"o número {num[2]} não está no perfil nem nas suas respostas"})
    ano_atual = __import__("datetime").date.today().year
    base_datas = datas(base)
    anos, meses = {a for a, _ in base_datas}, {(a, m) for a, m in base_datas if m}
    for ano, mes in datas(texto):
        if (mes and (ano, mes) in meses) or (not mes and ano in anos) or ano in (ano_atual, ano_atual + 1):
            continue
        valor = f"{mes:02d}/{ano}" if mes else str(ano)
        pontos.append({"tipo": "data", "valor": valor, "nivel": nivel_data, "mensagem": "data que não está no perfil"})
    return pontos


def conferir_carta(carta: str, perfil: str, vaga: str = "", empresa: str = "", respostas: dict | None = None) -> dict:
    """Veredito da carta: ok, conferir ou bloquear, com os pontos e a contagem de palavras."""
    respostas_txt = " ".join(str(v) for v in (respostas or {}).values())
    base = f"{perfil or ''}\n{respostas_txt}"
    pontos = _fatos_do_texto(carta, base, "bloquear")
    carta_n = norm(carta)
    n = contar_palavras(carta)
    minimo, maximo = PALAVRAS_CARTA
    if not minimo <= n <= maximo:
        pontos.append({"tipo": "tamanho", "valor": str(n), "nivel": "conferir",
                       "mensagem": f"a carta tem {n} palavras; o combinado é de {minimo} a {maximo}"})
    if empresa and norm(empresa) not in carta_n:
        pontos.append({"tipo": "empresa_da_vaga", "valor": empresa, "nivel": "conferir",
                       "mensagem": "a carta não cita a empresa da vaga: pode soar genérica"})
    reqs = [[t for t in termos(r) if len(t) > 2 and t not in VAZIAS_EN] for r in (requisitos_da_vaga(vaga) if vaga else [])]
    reqs = [r for r in reqs if r]
    if reqs and not any(sum(_contem(carta_n, t) for t in r) >= max(1, round(0.6 * len(r))) for r in reqs):
        pontos.append({"tipo": "requisito", "valor": "", "nivel": "conferir",
                       "mensagem": "a carta não cita nenhum requisito do anúncio: pode servir para qualquer vaga"})
    for e in EXPRESSOES_VAZIAS:
        if _contem(carta_n, norm(e)):
            pontos.append({"tipo": "vazia", "valor": e, "nivel": "conferir",
                           "mensagem": "expressão vazia; troque por um fato do perfil"})
    referencia = norm(f"{perfil} {vaga} {empresa} {respostas_txt}")
    for nome in _nomes_proprios(carta):
        if not _contem(referencia, norm(nome)):
            pontos.append({"tipo": "nome", "valor": nome, "nivel": "conferir",
                           "mensagem": "nome que não aparece no perfil nem na vaga (empresa, cargo ou lugar?)"})
    pior = max((NIVEIS[p["nivel"]] for p in pontos), default=0)
    return {"veredito": {0: "ok", 1: "conferir", 2: "bloquear"}[pior], "pontos": pontos, "palavras": n}


def conferir_respostas(itens: list[dict], perfil: str) -> list[dict]:
    """As respostas rascunhadas pela IA, cada uma com os pontos a conferir (fato fora do perfil)."""
    saida = []
    for i in itens:
        pontos = [] if i.get("da_pessoa") or not i.get("resposta") else _fatos_do_texto(i["resposta"], perfil or "", "conferir")
        saida.append({**i, "conferir": [p["mensagem"] + (f": {p['valor']}" if p["valor"] else "") for p in pontos]})
    return saida


def texto_carta(r: dict) -> str:
    rot = {"ok": "OK", "conferir": "CONFERIR", "bloquear": "BLOQUEAR"}
    linhas = [f"Veredito: {rot[r['veredito']]} · {r['palavras']} palavras"]
    linhas += [f"  [{p['nivel']}] {p['mensagem']}" + (f": {p['valor']}" if p["valor"] else "") for p in r["pontos"]]
    return "\n".join(linhas)


def texto(r: dict) -> str:
    rot = {"ok": "OK", "conferir": "CONFERIR", "bloquear": "BLOQUEAR"}
    linhas = [f"Veredito: {rot[r['veredito']]}"]
    if r["fatos"]:
        linhas.append("Fatos:")
        linhas += [f"  [{p['nivel']}] {p['mensagem']}: {p['valor']}" + (f" — \"{p['trecho']}\"" if p.get("trecho") and p["trecho"] != p["valor"] else "")
                   for p in r["fatos"]]
    else:
        linhas.append("Fatos: tudo bate com o perfil.")
    if r["requisitos"]:
        n = {s: sum(x["situacao"] == s for x in r["requisitos"]) for s in ("tem", "sustentado", "lacuna")}
        linhas.append(f"Requisitos da vaga ({len(r['requisitos'])}): {n['tem']} tem, {n['sustentado']} sustentado(s), {n['lacuna']} lacuna(s)")
        linhas += [f"  [{x['situacao']}] {x['texto']}" + (f" — {x['evidencia']}" if x["evidencia"] else "") for x in r["requisitos"]]
    if r["cobertura"] and r["cobertura"]["total"]:
        c = r["cobertura"]
        linhas.append(f"Palavras-chave da vaga no currículo: {c['achadas']} de {c['total']} ({round(100 * c['achadas'] / c['total'])}%)"
                      + (f"; faltando: {', '.join(c['faltando'])}" if c["faltando"] else ""))
    if r["ats"]:
        linhas.append("ATS:")
        linhas += [f"  [{p['nivel']}] {p['mensagem']}" + (f": {p['trecho']}" if p.get("trecho") else "") for p in r["ats"]]
    else:
        linhas.append("ATS: nada a apontar.")
    return "\n".join(linhas)


def perfil_padrao() -> Path:
    """O perfil apontado no config.json (padrão perfil.md)."""
    if str(RAIZ) not in sys.path:
        sys.path.insert(0, str(RAIZ))
    try:
        import filtros
        cfg = filtros.ler_config()
    except (ImportError, OSError, ValueError):
        cfg = {}
    return RAIZ / (cfg.get("perfil") or "perfil.md")


def vaga_do_banco(vid: str) -> dict:
    sys.path.insert(0, str(RAIZ / "dash"))
    import banco
    v = banco.obter(vid)
    if not v:
        raise ValueError(f"vaga {vid} não encontrada no dashboard")
    return v


def vaga_do_dashboard(vid: str) -> str:
    v = vaga_do_banco(vid)
    return "\n".join(str(x) for x in (v.get("titulo"), v.get("empresa"), v.get("local"), v.get("descricao")) if x)


def main(argv: list[str] | None = None) -> int:
    for fluxo in (sys.stdout, sys.stderr):
        try:
            fluxo.reconfigure(encoding="utf-8")
        except AttributeError:
            pass
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("arquivo", nargs="?", help="JSON do currículo")
    ap.add_argument("--carta", help="arquivo .txt de uma carta de apresentação (no lugar do currículo)")
    ap.add_argument("--respostas", help="JSON com as respostas da pessoa às perguntas da carta")
    ap.add_argument("--empresa", help="empresa da vaga (com --carta e --vaga)")
    ap.add_argument("--vaga", help="arquivo com o texto da vaga")
    ap.add_argument("--vaga-id", help="ID de uma vaga do dashboard")
    ap.add_argument("--perfil", help="perfil de carreira (padrão: o do config.json)")
    ap.add_argument("--json", action="store_true", help="resultado em JSON")
    args = ap.parse_args(argv)
    if bool(args.arquivo) == bool(args.carta):
        print("erro: informe o JSON do currículo ou --carta arquivo.txt", file=sys.stderr)
        return 2
    if args.carta:
        return _main_carta(args)
    try:
        cv = json.loads(Path(args.arquivo).read_text(encoding="utf-8"))
        perfil = Path(args.perfil) if args.perfil else perfil_padrao()
        perfil_txt = perfil.read_text(encoding="utf-8") if perfil.exists() else ""
        vaga = Path(args.vaga).read_text(encoding="utf-8") if args.vaga else vaga_do_dashboard(args.vaga_id) if args.vaga_id else None
    except (OSError, ValueError) as e:
        print(f"erro: {e}", file=sys.stderr)
        return 2
    if not perfil_txt:
        print(f"aviso: perfil não encontrado ({perfil}); todos os fatos ficam para conferir", file=sys.stderr)
    r = conferir(cv, perfil_txt, vaga, Path(args.arquivo).with_suffix(".docx"))
    print(json.dumps(r, ensure_ascii=False, indent=1) if args.json else texto(r))
    return {"ok": 0, "conferir": 1, "bloquear": 3}[r["veredito"]]


def _main_carta(args) -> int:
    try:
        carta = Path(args.carta).read_text(encoding="utf-8")
        perfil = Path(args.perfil) if args.perfil else perfil_padrao()
        perfil_txt = perfil.read_text(encoding="utf-8") if perfil.exists() else ""
        respostas = json.loads(Path(args.respostas).read_text(encoding="utf-8")) if args.respostas else {}
        if args.vaga_id:
            v = vaga_do_banco(args.vaga_id)
            vaga, empresa = vaga_do_dashboard(args.vaga_id), v.get("empresa") or ""
        else:
            vaga = Path(args.vaga).read_text(encoding="utf-8") if args.vaga else ""
            empresa = args.empresa or ""
    except (OSError, ValueError) as e:
        print(f"erro: {e}", file=sys.stderr)
        return 2
    r = conferir_carta(carta, perfil_txt, vaga, empresa, respostas if isinstance(respostas, dict) else {})
    print(json.dumps(r, ensure_ascii=False, indent=1) if args.json else texto_carta(r))
    return {"ok": 0, "conferir": 1, "bloquear": 3}[r["veredito"]]


if __name__ == "__main__":
    sys.exit(main())
